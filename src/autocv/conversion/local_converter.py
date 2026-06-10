from pathlib import Path

from autocv.conversion.contracts import ConversionRequest, ConversionResponse, DocumentFormat


class LocalDocumentConverter:
    def convert(self, request: ConversionRequest) -> ConversionResponse:
        try:
            if request.source_format == DocumentFormat.DOCX and request.target_format == DocumentFormat.PDF:
                _docx_to_pdf(request.source_path, request.output_path)
            elif request.source_format == DocumentFormat.XLSX and request.target_format == DocumentFormat.PDF:
                _xlsx_to_pdf(request.source_path, request.output_path)
            elif request.source_format == DocumentFormat.PDF and request.target_format == DocumentFormat.DOCX:
                _pdf_to_docx(request.source_path, request.output_path)
            elif request.source_format == DocumentFormat.PDF and request.target_format == DocumentFormat.XLSX:
                _pdf_to_xlsx(request.source_path, request.output_path)
            else:
                return ConversionResponse(
                    output_path=request.output_path,
                    source="local_converter",
                    available=False,
                    message="Conversion non supportée en V1.",
                )
        except Exception as exc:
            return ConversionResponse(
                output_path=request.output_path,
                source="local_converter",
                available=False,
                message=_friendly_conversion_error(exc),
            )

        return ConversionResponse(
            output_path=request.output_path,
            source="local_converter",
            available=True,
            message="Conversion terminée.",
        )

def _friendly_conversion_error(error: Exception) -> str:
    module_name = type(error).__module__
    text = str(error)
    if isinstance(error, ModuleNotFoundError) and "win32com" in text:
        return "Conversion impossible: pywin32 ou Microsoft Word/Excel est indisponible."
    if "pywintypes" in module_name or "win32com" in module_name:
        return (
            "Conversion impossible: Microsoft Word/Excel est indisponible "
            "ou le document est verrouille."
        )
    return f"Conversion impossible: {text}"


def _docx_to_pdf(source_path: Path, output_path: Path) -> None:
    import win32com.client

    output_path.parent.mkdir(parents=True, exist_ok=True)
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    document = None
    try:
        document = word.Documents.Open(str(source_path.resolve()))
        document.SaveAs(str(output_path.resolve()), FileFormat=17)
    finally:
        if document is not None:
            document.Close(False)
        word.Quit()


def _xlsx_to_pdf(source_path: Path, output_path: Path) -> None:
    import win32com.client

    output_path.parent.mkdir(parents=True, exist_ok=True)
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    workbook = None
    try:
        workbook = excel.Workbooks.Open(str(source_path.resolve()))
        workbook.ExportAsFixedFormat(0, str(output_path.resolve()))
    finally:
        if workbook is not None:
            workbook.Close(False)
        excel.Quit()


def _pdf_to_docx(source_path: Path, output_path: Path) -> None:
    from docx import Document
    from pypdf import PdfReader

    output_path.parent.mkdir(parents=True, exist_ok=True)
    reader = PdfReader(str(source_path))
    document = Document()
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if index > 1:
            document.add_page_break()
        for line in text.splitlines():
            document.add_paragraph(line)
    document.save(str(output_path))


def _pdf_to_xlsx(source_path: Path, output_path: Path) -> None:
    from openpyxl import Workbook
    from pypdf import PdfReader

    output_path.parent.mkdir(parents=True, exist_ok=True)
    reader = PdfReader(str(source_path))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PDF"
    row_index = 1
    for page in reader.pages:
        text = page.extract_text() or ""
        for line in text.splitlines():
            values = [part for part in line.split() if part]
            if values:
                for column_index, value in enumerate(values, start=1):
                    sheet.cell(row=row_index, column=column_index).value = value
                row_index += 1
    workbook.save(str(output_path))
